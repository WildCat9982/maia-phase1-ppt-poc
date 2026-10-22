"""Create a small sample deck for demos and tests: python scripts/make_sample_deck.py [out.pptx]"""

import sys

from pptx import Presentation
from pptx.util import Inches


def build(path: str) -> None:
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Project Falcon: Quarterly Review"
    s.placeholders[1].text = "Programme Office · Q3"

    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Progress this quarter"
    tf = s.placeholders[1].text_frame
    tf.text = "Design review completed on schedule"
    for line in ["Two of three integration tests passed", "Vendor delivery slipped by three weeks"]:
        tf.add_paragraph().text = line

    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Budget status"
    rows = [("Item", "Planned", "Actual"), ("Hardware", "1.2M", "1.3M"), ("Services", "0.8M", "0.7M")]
    table = s.shapes.add_table(3, 3, Inches(1), Inches(2), Inches(8), Inches(1.5)).table
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            table.cell(r, c).text = val

    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Decisions needed"
    tf = s.placeholders[1].text_frame
    tf.text = "Approve a three-week schedule extension"
    tf.add_paragraph().text = "Contact the project lead at falcon.lead@example.gov.sg or 9123 4567"
    s.notes_slide.notes_text_frame.text = "Existing note: stress the extension is low risk."

    prs.save(path)


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "sample_deck.pptx"
    build(out)
    print(f"Wrote {out}")
