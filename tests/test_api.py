"""End-to-end API tests with the mock LLM: pytest -q"""

import io
import json
import os
import sys
import zipfile
from pathlib import Path

os.environ["LLM_PROVIDER"] = "mock"
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from pptx import Presentation  # noqa: E402

from backend.main import app  # noqa: E402
from backend.masking import mask  # noqa: E402
from make_sample_deck import build  # noqa: E402

MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
client = TestClient(app)


@pytest.fixture(scope="module")
def sample(tmp_path_factory) -> bytes:
    path = tmp_path_factory.mktemp("d") / "sample.pptx"
    build(str(path))
    return path.read_bytes()


@pytest.fixture(scope="module")
def deck(sample) -> dict:
    resp = client.post("/api/ppt/extract", files={"file": ("sample.pptx", sample, MIME)})
    assert resp.status_code == 200
    return resp.json()


def test_health():
    assert client.get("/health").json()["provider"] == "mock"


def test_extract(deck):
    assert deck["slide_count"] == 4
    assert deck["slides"][1]["title"] == "Progress this quarter"
    assert deck["slides"][2]["tables"][0][1] == ["Hardware", "1.2M", "1.3M"]
    assert deck["slides"][3]["notes"].startswith("Existing note")


def test_rejects_non_pptx():
    resp = client.post("/api/ppt/extract", files={"file": ("x.txt", b"hello", "text/plain")})
    assert resp.status_code == 415


def test_summarise_citations_are_valid(deck):
    resp = client.post("/api/ppt/summarise", json=deck)
    assert resp.status_code == 200
    summary = resp.json()["summary"]
    cited = {n for p in summary["key_points"] for n in p["slides"]}
    assert cited and cited <= {1, 2, 3, 4}


def test_notes_keep_existing_and_apply(deck, sample):
    resp = client.post("/api/ppt/notes", json={"deck": deck, "seconds_per_slide": 30})
    notes = resp.json()["notes"]
    assert len(notes) == 4
    assert notes[3]["skipped"] is True  # slide 4 already had notes
    payload = {str(n["index"]): n["notes"] for n in notes}
    out = client.post("/api/ppt/notes/apply", files={"file": ("sample.pptx", sample, MIME)},
                      data={"notes_json": json.dumps(payload)})
    assert out.status_code == 200
    prs = Presentation(io.BytesIO(out.content))
    assert prs.slides[0].notes_slide.notes_text_frame.text == notes[0]["notes"]


def test_outline_and_render_default_template():
    resp = client.post("/api/ppt/outline", json={"prompt": "Brief on Project Falcon", "slide_count": 4})
    body = resp.json()
    assert len(body["outline"]["slides"]) == 4
    assert body["quality"]["slides"] == 5
    out = client.post("/api/ppt/render", data={"outline_json": json.dumps(body["outline"])})
    prs = Presentation(io.BytesIO(out.content))
    assert len(prs.slides) == 5
    assert prs.slides[0].shapes.title.text == body["outline"]["title"]


def test_render_on_potx_template(sample):
    # Turn the sample deck into a .potx to prove templates load and old slides are removed.
    src = zipfile.ZipFile(io.BytesIO(sample))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"presentationml.presentation.main+xml", b"presentationml.template.main+xml")
            dst.writestr(item, data)
    outline = {"title": "Templated deck", "slides": [
        {"layout": "section", "title": "Part one"},
        {"layout": "bullets", "title": "Points", "bullets": ["a", "b"], "notes": "say this"}]}
    out = client.post("/api/ppt/render", data={"outline_json": json.dumps(outline)},
                      files={"template": ("brand.potx", buf.getvalue(), "application/octet-stream")})
    assert out.status_code == 200
    prs = Presentation(io.BytesIO(out.content))
    assert [s.shapes.title.text for s in prs.slides] == ["Templated deck", "Part one", "Points"]
    assert prs.slides[2].notes_slide.notes_text_frame.text == "say this"


def test_masking():
    text = "Call 9123 4567 or +65 6123-4567, mail a.b@example.gov.sg, NRIC S1234567D"
    masked = mask(text)
    assert "9123" not in masked and "6123" not in masked
    assert "[EMAIL]" in masked and "[NRIC]" in masked
