"""Reading and writing PowerPoint files with python-pptx.

In the production design the VSTO plugin extracts the deck and inserts slides;
in this POC the backend does both from uploaded files, using the same deck JSON.
"""

import io
import zipfile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.util import Pt

from .schemas import Deck, Outline, QualityReport, SlideContent

_TEMPLATE_CT = "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
_PRESENTATION_CT = "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"


# ---- Opening files -----------------------------------------------------------

def open_presentation(data: bytes | None) -> Presentation:
    """Open .pptx bytes, or .potx bytes (python-pptx refuses templates as-is)."""
    if data is None:
        return Presentation()
    return Presentation(io.BytesIO(_potx_to_pptx(data)))


def _potx_to_pptx(data: bytes) -> bytes:
    """A .potx differs from a .pptx only in one content type; rewrite it."""
    src = zipfile.ZipFile(io.BytesIO(data))
    types = src.read("[Content_Types].xml").decode("utf-8")
    if _TEMPLATE_CT not in types:
        return data
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            content = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                content = types.replace(_TEMPLATE_CT, _PRESENTATION_CT).encode("utf-8")
            dst.writestr(item, content)
    return out.getvalue()


def to_bytes(prs: Presentation) -> bytes:
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ---- Extraction (deck -> JSON) -------------------------------------------------

def _iter_shapes(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def extract(data: bytes) -> Deck:
    prs = open_presentation(data)
    slides: list[SlideContent] = []
    for i, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None else ""
        body: list[str] = []
        tables: list[list[list[str]]] = []
        has_images = False
        for shape in _iter_shapes(slide.shapes):
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                has_images = True
            if getattr(shape, "has_table", False) and shape.has_table:
                tables.append([[cell.text.strip() for cell in row.cells] for row in shape.table.rows])
            elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs).strip()
                    if text:
                        body.append(("  " * para.level) + text)
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        slides.append(SlideContent(
            index=i, layout=slide.slide_layout.name, title=title, body=body,
            tables=tables, notes=notes, has_images=has_images,
        ))
    return Deck(slide_count=len(slides), slides=slides)


def deck_as_text(deck: Deck) -> str:
    """Plain-text form of the deck sent to the model, one block per slide."""
    parts = []
    for s in deck.slides:
        lines = [f"[Slide {s.index}] {s.title or '(no title)'}"]
        lines += [f"- {b}" for b in s.body]
        for t, table in enumerate(s.tables, start=1):
            lines.append(f"Table {t}:")
            lines += [" | ".join(row) for row in table]
        if s.has_images:
            lines.append("(slide contains images not shown here)")
        parts.append("\n".join(lines))
    return "\n\n".join(parts)


# ---- Writing speaker notes -----------------------------------------------------

def apply_notes(data: bytes, notes: dict[int, str]) -> bytes:
    """Return a copy of the deck with notes written to the given slide numbers."""
    prs = open_presentation(data)
    for i, slide in enumerate(prs.slides, start=1):
        if i in notes:
            slide.notes_slide.notes_text_frame.text = notes[i]
    return to_bytes(prs)


# ---- Rendering an outline onto a template -------------------------------------

_LAYOUT_NAMES = {
    "title": ["title slide"],
    "section": ["section header", "section title", "section"],
    "bullets": ["title and content", "title, content", "content"],
}
_LAYOUT_FALLBACK = {"title": 0, "bullets": 1, "section": 2}


def _pick_layout(prs: Presentation, kind: str):
    layouts = list(prs.slide_layouts)
    for wanted in _LAYOUT_NAMES[kind]:
        for layout in layouts:
            if layout.name.strip().lower() == wanted:
                return layout
    for wanted in _LAYOUT_NAMES[kind]:
        for layout in layouts:
            if wanted in layout.name.strip().lower():
                return layout
    return layouts[min(_LAYOUT_FALLBACK[kind], len(layouts) - 1)]


def _remove_all_slides(prs: Presentation) -> None:
    id_list = prs.slides._sldIdLst  # python-pptx has no public delete API
    for sld_id in list(id_list):
        prs.part.drop_rel(sld_id.rId)
        id_list.remove(sld_id)


def _body_placeholder(slide):
    for ph in slide.placeholders:
        kind = ph.placeholder_format.type
        if kind in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.SUBTITLE):
            return ph
    return None


def render(outline: Outline, template: bytes | None) -> bytes:
    """Build a .pptx from the outline using the template's own layouts."""
    prs = open_presentation(template)
    _remove_all_slides(prs)

    def add(kind: str, title: str, bullets: list[str], notes: str):
        slide = prs.slides.add_slide(_pick_layout(prs, kind))
        if slide.shapes.title is not None:
            slide.shapes.title.text = title
        body = _body_placeholder(slide)
        if body is not None:
            if bullets:
                tf = body.text_frame
                tf.text = bullets[0]
                for b in bullets[1:]:
                    tf.add_paragraph().text = b
            else:
                body._element.getparent().remove(body._element)  # no empty placeholder
        elif bullets:  # layout has no body placeholder: add a plain text box
            box = slide.shapes.add_textbox(Pt(60), Pt(140), prs.slide_width - Pt(120), Pt(300))
            box.text_frame.text = "\n".join(bullets)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    add("title", outline.title, [outline.subtitle] if outline.subtitle else [], "")
    for s in outline.slides:
        if s.layout == "title":  # only one title slide; treat extras as sections
            add("section", s.title, s.bullets, s.notes)
        else:
            add(s.layout, s.title, s.bullets, s.notes)
    return to_bytes(prs)


def check_outline(outline: Outline) -> QualityReport:
    """Cheap proxies for the 'overflowing text' metric in the proposal."""
    over, long_, empty = [], [], []
    for n, s in enumerate(outline.slides, start=2):  # slide 1 is the title slide
        if len(s.bullets) > 5:
            over.append(n)
        if any(len(b.split()) > 12 for b in s.bullets):
            long_.append(n)
        if s.layout == "bullets" and not s.bullets:
            empty.append(n)
    return QualityReport(slides=len(outline.slides) + 1, bullets_over_limit=over,
                         long_bullets=long_, empty_slides=empty)
